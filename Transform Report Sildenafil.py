import os
import re
from openpyxl import load_workbook
import pandas as pd
from datetime import datetime
import json
import subprocess

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DIR_SOURCE = os.path.join(BASE_DIR, "source")
DIR_IMAGE = os.path.join(BASE_DIR, "image")
DIR_OUTPUT = os.path.join(BASE_DIR, "output")
FILE_MAPPING = os.path.join(BASE_DIR, "mapping_surveyor.xlsx")
FILE_LOG = os.path.join(BASE_DIR, "log_proses_survey.xlsx")
FILE_HTML_PETA = os.path.join(BASE_DIR, "peta_sebaran_sildenafil.html")

for direktori in [DIR_SOURCE, DIR_IMAGE, DIR_OUTPUT]:
    if not os.path.exists(direktori):
        os.makedirs(direktori)

if not os.path.exists(FILE_MAPPING):
    df_template = pd.DataFrame(columns=["Surveyor ID", "Sales Name", "Position", "Region", "ASM"])
    df_template.to_excel(FILE_MAPPING, index=False)
    print(f"[INFO] File '{os.path.basename(FILE_MAPPING)}' baru saja dibuat.")
    print("-> Silakan isi data Surveyor kamu di file Excel tersebut dahulu, lalu jalankan ulang script ini.")
    exit()

print("-> Memuat data master dari Excel mapping...")
df_mapping_excel = pd.read_excel(FILE_MAPPING)

MASTER_SURVEYOR = {}
for idx, row in df_mapping_excel.iterrows():
    if pd.isna(row["Surveyor ID"]):
        continue
    sid = str(row["Surveyor ID"]).strip()
    MASTER_SURVEYOR[sid] = {
        "Position": row["Position"] if not pd.isna(row["Position"]) else "-",
        "Region": row["Region"] if not pd.isna(row["Region"]) else "-",
        "ASM": row["ASM"] if not pd.isna(row["ASM"]) else "-"
    }

files_in_source = [f for f in os.listdir(DIR_SOURCE) if f.endswith('.xlsx') and not f.startswith('~$')]

if not files_in_source:
    print(f"[ERROR] Tidak ditemukan file .xlsx di: {DIR_SOURCE}")
    exit()

print(f"Ditemukan {len(files_in_source)} file Excel untuk diproses.\n")

outlet_aggregated_data = {}
kolom_headers = []
list_log_detail = []

total_foto_diekstrak = 0
total_baris_digabung = 0
total_foto_gagal = 0
total_tanpa_foto = 0

for nama_file_excel in files_in_source:
    path_excel_input = os.path.join(DIR_SOURCE, nama_file_excel)
    print(f"==================================================")
    print(f"-> Membaca file: {nama_file_excel}")
    print(f"==================================================")

    wb_image = load_workbook(path_excel_input, data_only=True)
    ws_image = wb_image.active

    if not kolom_headers:
        kolom_headers = [ws_image.cell(row=2, column=col).value for col in range(1, ws_image.max_column + 1)]

    lokasi_col_idx = None
    for col_idx, header in enumerate(kolom_headers, start=1):
        if header:
            h_str = str(header).lower().strip()
            if ('survey' in h_str and ('loc' in h_str or 'coord' in h_str or 'koordinat' in h_str or 'lat' in h_str)) or (h_str in ['survey location', 'lokasi survey', 'survey_location']):
                lokasi_col_idx = col_idx
                break
    
    if lokasi_col_idx is None:
        for col_idx, header in enumerate(kolom_headers, start=1):
            if header:
                h_str = str(header).lower().strip()
                if 'survey' in h_str and 'visit' not in h_str:
                    lokasi_col_idx = col_idx
                    break

    gambar_per_baris = {}
    for img in ws_image._images:
        baris_gambar = img.anchor._from.row + 1
        gambar_per_baris[baris_gambar] = img

    for row in range(3, ws_image.max_row + 1):
        timestamp_raw = ws_image.cell(row=row, column=1).value
        surveyor_id_raw = ws_image.cell(row=row, column=2).value
        surveyor = ws_image.cell(row=row, column=3).value
        surveyee_id = ws_image.cell(row=row, column=4).value
        surveyee = ws_image.cell(row=row, column=5).value

        if not surveyor or not surveyee_id or not surveyee:
            continue
            
        timestamp_str = str(timestamp_raw).strip() if timestamp_raw else f"Tanpa_Timestamp_{row}"
        surveyor_id_str = str(surveyor_id_raw).strip() if surveyor_id_raw else ""
        surveyor_str = str(surveyor).strip()
        surveyee_id_str = str(surveyee_id).strip()
        surveyee_str = str(surveyee).strip()

        row_values = [ws_image.cell(row=row, column=col).value for col in range(1, ws_image.max_column + 1)]
        
        lat, lon = None, None
        def parse_and_fix_latlon(val_str):
            if not val_str:
                return None, None
            parts = re.split(r'[,\s]+', str(val_str).strip())
            floats = []
            for p in parts:
                try:
                    floats.append(float(p))
                except ValueError:
                    continue
            if len(floats) >= 2:
                v1, v2 = floats[0], floats[1]
                if (95.0 <= abs(v1) <= 141.0) and (-11.0 <= v2 <= 6.0):
                    return v2, v1
                elif (-11.0 <= v1 <= 6.0) and (95.0 <= abs(v2) <= 141.0):
                    return v1, v2
            return None, None

        if lokasi_col_idx:
            loc_val = ws_image.cell(row=row, column=lokasi_col_idx).value
            lat, lon = parse_and_fix_latlon(loc_val)

        row_qty = 0
        for idx_col, val_col in enumerate(row_values):
            header_name = str(kolom_headers[idx_col]).lower() if idx_col < len(kolom_headers) else ""
            if any(k in header_name for k in ['qty', 'jumlah', 'total', 'pack', 'pcs', 'quantity']):
                if pd.notna(val_col):
                    nums = re.findall(r'\d+', str(val_col).replace(',', '.'))
                    if nums:
                        try:
                            row_qty += float(nums[0])
                        except:
                            pass

        path_simpan_relatif_list = []
        status_foto_log = "-"
        keterangan_foto_log = "-"

        if row in gambar_per_baris:
            surveyor_clean = re.sub(r'[\/*?:"<>|]', '', surveyor_str).strip(' .')
            surveyee_clean = re.sub(r'[\/*?:"<>|]', '', surveyee_str).strip(' .')
            surveyee_id_clean = re.sub(r'[\/*?:"<>|]', '', surveyee_id_str).strip(' .')
            timestamp_clean = re.sub(r'[\s:]', '_', timestamp_str).strip(' .')

            folder_tujuan_foto = os.path.join(DIR_IMAGE, surveyor_clean, surveyee_clean)
            if not os.path.exists(folder_tujuan_foto):
                os.makedirs(folder_tujuan_foto)

            nama_file_foto = f"{surveyee_id_clean}_{surveyee_clean}_{timestamp_clean}_r{row}.jpg"
            path_simpan_foto = os.path.join(folder_tujuan_foto, nama_file_foto)
            path_simpan_relatif = f"image/{surveyor_clean}/{surveyee_clean}/{nama_file_foto}".replace(os.sep, "/")

            img_obj = gambar_per_baris[row]
            try:
                if hasattr(img_obj, 'image'):
                    img_obj.image.save(path_simpan_foto)
                elif hasattr(img_obj, 'ref') and hasattr(img_obj.ref, 'read'):
                    img_obj.ref.seek(0)
                    with open(path_simpan_foto, 'wb') as f:
                        f.write(img_obj.ref.read())
                elif hasattr(img_obj, 'fp') and hasattr(img_obj.fp, 'read'):
                    img_obj.fp.seek(0)
                    with open(path_simpan_foto, 'wb') as f:
                        f.write(img_obj.fp.read())
                
                total_foto_diekstrak += 1
                path_simpan_relatif_list.append(path_simpan_relatif)
                status_foto_log = "SUKSES"
                keterangan_foto_log = path_simpan_foto
            except Exception as e:
                total_foto_gagal += 1
                status_foto_log = "GAGAL"
                keterangan_foto_log = str(e)
        else:
            total_tanpa_foto += 1
            status_foto_log = "TIDAK ADA FOTO"
            keterangan_foto_log = "Tidak ada objek gambar pada baris ini"

        outlet_key = (surveyor_id_str, surveyor_str, surveyee_id_str, surveyee_str)

        if outlet_key in outlet_aggregated_data:
            total_baris_digabung += 1
            existing_qty = outlet_aggregated_data[outlet_key]["qty_numeric"]
            new_qty = existing_qty + row_qty
            outlet_aggregated_data[outlet_key]["qty_numeric"] = new_qty
            outlet_aggregated_data[outlet_key]["qty"] = int(new_qty) if new_qty.is_integer() else new_qty
            
            if path_simpan_relatif_list:
                outlet_aggregated_data[outlet_key]["foto_list"].extend(path_simpan_relatif_list)
            
            outlet_aggregated_data[outlet_key]["timestamp"] = f"{outlet_aggregated_data[outlet_key]['timestamp']}, {timestamp_str}"

            list_log_detail.append({
                "Waktu Eksekusi": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "Nama File Source": nama_file_excel,
                "Nomor Baris": row,
                "Timestamp Data": timestamp_str,
                "Surveyor ID": surveyor_id_str,
                "Surveyor": surveyor_str,
                "Surveyee ID": surveyee_id_str,
                "Surveyee / Outlet": surveyee_str,
                "Status Baris": "DIGABUNG (QTY Dijumlah)",
                "Status Foto": status_foto_log,
                "Path / Keterangan": f"Outlet sudah ada, Qty digabung jadi {outlet_aggregated_data[outlet_key]['qty']}"
            })
        else:
            outlet_aggregated_data[outlet_key] = {
                "surveyor_id": surveyor_id_str,
                "surveyor": surveyor_str,
                "surveyee_id": surveyee_id_str,
                "surveyee": surveyee_str,
                "timestamp": timestamp_str,
                "lat": lat,
                "lon": lon,
                "qty_numeric": row_qty,
                "qty": int(row_qty) if row_qty.is_integer() else (row_qty if row_qty > 0 else "-"),
                "foto_list": path_simpan_relatif_list,
                "row_values": row_values
            }
            list_log_detail.append({
                "Waktu Eksekusi": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "Nama File Source": nama_file_excel,
                "Nomor Baris": row,
                "Timestamp Data": timestamp_str,
                "Surveyor ID": surveyor_id_str,
                "Surveyor": surveyor_str,
                "Surveyee ID": surveyee_id_str,
                "Surveyee / Outlet": surveyee_str,
                "Status Baris": "BARU (Tersimpan)",
                "Status Foto": status_foto_log,
                "Path / Keterangan": keterangan_foto_log
            })

print(f"Total outlet unik setelah deduplikasi & agregasi Qty: {len(outlet_aggregated_data)} outlet")
print("=== Menghasilkan Dashboard Peta GIS Indonesia (Mobile Responsive) ===")

semua_data_gabungan = []
list_data_peta = []

for k, v in outlet_aggregated_data.items():
    semua_data_gabungan.append(v["row_values"])
    sid = v["surveyor_id"]
    reg = MASTER_SURVEYOR.get(sid, {}).get("Region", "-")
    asm = MASTER_SURVEYOR.get(sid, {}).get("ASM", "-")

    if v["lat"] is not None and v["lon"] is not None:
        list_data_peta.append({
            "lat": v["lat"],
            "lon": v["lon"],
            "timestamp": v["timestamp"],
            "surveyor": v["surveyor"],
            "surveyor_id": v["surveyor_id"],
            "surveyee": v["surveyee"],
            "surveyee_id": v["surveyee_id"],
            "region": reg,
            "asm": asm,
            "qty": v["qty"],
            "foto_list": v["foto_list"]
        })

if semua_data_gabungan:
    df_clean = pd.DataFrame(semua_data_gabungan, columns=kolom_headers)

    list_lat_excel = []
    list_lon_excel = []
    def parse_and_fix_latlon(val_str):
        if not val_str:
            return None, None
        parts = re.split(r'[,\s]+', str(val_str).strip())
        floats = []
        for p in parts:
            try:
                floats.append(float(p))
            except ValueError:
                continue
        if len(floats) >= 2:
            v1, v2 = floats[0], floats[1]
            if (95.0 <= abs(v1) <= 141.0) and (-11.0 <= v2 <= 6.0):
                return v2, v1
            elif (-11.0 <= v1 <= 6.0) and (95.0 <= abs(v2) <= 141.0):
                return v1, v2
        return None, None

    for idx, row in df_clean.iterrows():
        loc_val = row.iloc[lokasi_col_idx - 1] if lokasi_col_idx and lokasi_col_idx <= len(row) else None
        lt, ln = parse_and_fix_latlon(loc_val)
        list_lat_excel.append(lt if lt is not None else "-")
        list_lon_excel.append(ln if ln is not None else "-")

    df_clean["Latitude"] = list_lat_excel
    df_clean["Longitude"] = list_lon_excel

    list_position = []
    list_region = []
    list_asm = []

    for idx, row in df_clean.iterrows():
        sid = str(row['Surveyor ID']).strip() if 'Surveyor ID' in df_clean.columns else ""
        if sid in MASTER_SURVEYOR:
            list_position.append(MASTER_SURVEYOR[sid]["Position"])
            list_region.append(MASTER_SURVEYOR[sid]["Region"])
            list_asm.append(MASTER_SURVEYOR[sid]["ASM"])
        else:
            list_position.append("-")
            list_region.append("-")
            list_asm.append("-")

    df_clean["Position"] = list_position
    df_clean["Region"] = list_region
    df_clean["ASM"] = list_asm

    path_excel_output = os.path.join(DIR_OUTPUT, "all_data_combined_clean.xlsx")
    
    try:
        df_clean.to_excel(path_excel_output, index=False)
        df_log = pd.DataFrame(list_log_detail)
        df_log.to_excel(FILE_LOG, index=False)
        data_json = json.dumps(list_data_peta)

        multilayer_dashboard_html = f"""<!DOCTYPE html>
<html lang="id">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0, maximum-scale=1.0, user-scalable=no" />
    <title>Dashboard Distribusi Sildenafil Package</title>
    <link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css" />
    <script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
    <style>
        * {{ box-sizing: border-box; -webkit-tap-highlight-color: transparent; }}
        html, body {{ height: 100%; margin: 0; padding: 0; font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; background-color: #f1f5f9; color: #1e293b; overflow: hidden; }}
        header {{ height: 55px; background-color: #ffffff; color: #0f172a; padding: 0 15px; display: flex; justify-content: space-between; align-items: center; border-bottom: 1px solid #cbd5e1; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.05); z-index: 2000; position: relative; }}
        header h1 {{ margin: 0; font-size: 15px; font-weight: 600; display: flex; align-items: center; gap: 8px; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
        .header-badge {{ background: #2563eb; color: white; padding: 4px 10px; border-radius: 6px; font-size: 11px; font-weight: 600; white-space: nowrap; }}
        .main-container {{ display: flex; height: calc(100vh - 55px); width: 100vw; position: relative; }}
        .sidebar {{ width: 380px; background-color: #ffffff; border-right: 1px solid #cbd5e1; display: flex; flex-direction: column; padding: 15px; gap: 10px; z-index: 1000; box-shadow: 4px 0 15px rgba(0,0,0,0.05); overflow-y: auto; }}
        .filter-group {{ display: flex; flex-direction: column; gap: 3px; }}
        .filter-group label {{ font-size: 11px; font-weight: 600; color: #64748b; }}
        .filter-group select {{ padding: 7px 10px; border-radius: 6px; border: 1px solid #cbd5e1; background-color: #f8fafc; color: #1e293b; font-size: 12px; outline: none; }}
        .stats-box {{ background: linear-gradient(135deg, #2563eb 0%, #1d4ed8 100%); padding: 10px; border-radius: 8px; color: white; box-shadow: 0 4px 12px rgba(37, 99, 235, 0.2); }}
        .stats-box h3 {{ margin: 0 0 2px 0; font-size: 10px; text-transform: uppercase; letter-spacing: 0.5px; opacity: 0.9; }}
        .stats-box .val {{ font-size: 18px; font-weight: bold; }}
        .outlet-list {{ flex: 1; min-height: 150px; overflow-y: auto; display: flex; flex-direction: column; gap: 6px; padding-right: 2px; }}
        .outlet-item {{ background-color: #f8fafc; border: 1px solid #cbd5e1; padding: 8px 10px; border-radius: 6px; cursor: pointer; transition: all 0.2s ease; }}
        .outlet-item:hover {{ border-color: #2563eb; background-color: #eff6ff; }}
        .outlet-item.active {{ border-color: #2563eb; background-color: #dbeafe; }}
        .outlet-name {{ font-weight: 600; font-size: 12px; color: #0f172a; margin-bottom: 2px; }}
        .outlet-info {{ font-size: 11px; color: #64748b; display: flex; justify-content: space-between; }}
        #map {{ flex: 1; height: 100%; width: 100%; z-index: 1; }}

        @media (max-width: 768px) {{
            .main-container {{ flex-direction: column-reverse; height: calc(100vh - 55px); }}
            .sidebar {{ width: 100%; height: 45%; max-height: 250px; padding: 10px; gap: 6px; border-right: none; border-top: 1px solid #cbd5e1; }}
            #map {{ height: 55% !important; width: 100% !important; }}
            header h1 {{ font-size: 13px; }}
        }}
        ::-webkit-scrollbar {{ width: 5px; }}
        ::-webkit-scrollbar-track {{ background: #f1f5f9; }}
        ::-webkit-scrollbar-thumb {{ background: #cbd5e1; border-radius: 3px; }}
    </style>
</head>
<body>
    <header>
        <h1>🗺️ Distribusi Sildenafil Package</h1>
        <div id="counterBadge" class="header-badge">Memuat...</div>
    </header>
    <div class="main-container">
        <div class="sidebar">
            <div class="filter-group">
                <label for="regionFilter">Filter Region:</label>
                <select id="regionFilter" onchange="onRegionChange()">
                    <option value="ALL">-- Semua Region --</option>
                </select>
            </div>
            <div class="filter-group">
                <label for="asmFilter">Filter ASM:</label>
                <select id="asmFilter" onchange="onAsmChange()">
                    <option value="ALL">-- Semua ASM --</option>
                </select>
            </div>
            <div class="filter-group">
                <label for="surveyorFilter">Filter Surveyor:</label>
                <select id="surveyorFilter" onchange="onSurveyorChange()">
                    <option value="ALL">-- Semua Surveyor --</option>
                </select>
            </div>
            <div class="filter-group">
                <label for="outletFilter">Filter Outlet (Surveyee):</label>
                <select id="outletFilter" onchange="onOutletChange()">
                    <option value="ALL">-- Semua Outlet --</option>
                </select>
            </div>
            <div class="stats-box">
                <h3>Total Distribusi Sildenafil Package</h3>
                <div id="totalOutletStat" class="val">0</div>
            </div>
            <div style="font-size: 11px; font-weight: 600; color: #64748b;">Daftar Outlet Terfilter:</div>
            <div id="outletListContainer" class="outlet-list"></div>
        </div>
        <div id="map"></div>
    </div>
    <script>
        const surveyData = {data_json};
        let surveyorColorMap = {{}};
        let map, markersLayer;

        function initDashboard() {{
            populateRegionDropdown();
            populateAsmDropdown(surveyData);
            populateSurveyorDropdown(surveyData);
            populateOutletDropdown(surveyData);
            initMap();
            renderDashboard('ALL', 'ALL', 'ALL', 'ALL');
        }}

        function populateRegionDropdown() {{
            const select = document.getElementById('regionFilter');
            const regions = [...new Set(surveyData.map(item => item.region))].sort();
            regions.forEach(r => {{
                if(r && r !== '-') {{
                    const opt = document.createElement('option');
                    opt.value = r;
                    opt.textContent = r;
                    select.appendChild(opt);
                }}
            }});
        }}

        function populateAsmDropdown(dataList) {{
            const select = document.getElementById('asmFilter');
            const currentVal = select.value;
            select.innerHTML = '<option value="ALL">-- Semua ASM --</option>';
            const asms = [...new Set(dataList.map(item => item.asm))].sort();
            asms.forEach(a => {{
                if(a && a !== '-') {{
                    const opt = document.createElement('option');
                    opt.value = a;
                    opt.textContent = a;
                    select.appendChild(opt);
                }}
            }});
            select.value = asms.includes(currentVal) ? currentVal : 'ALL';
        }}

        function populateSurveyorDropdown(dataList) {{
            const select = document.getElementById('surveyorFilter');
            const currentVal = select.value;
            select.innerHTML = '<option value="ALL">-- Semua Surveyor --</option>';
            const surveyors = [...new Set(dataList.map(item => item.surveyor))].sort();
            surveyors.forEach(s => {{
                if(s && s !== '-') {{
                    const opt = document.createElement('option');
                    opt.value = s;
                    opt.textContent = s;
                    select.appendChild(opt);
                }}
            }});
            select.value = surveyors.includes(currentVal) ? currentVal : 'ALL';

            const palette = ['#2563eb', '#dc2626', '#059669', '#d97706', '#7c3aed', '#db2777', '#0891b2', '#4f46e5', '#ca8a04', '#65a30d'];
            surveyors.forEach((s, idx) => {{ surveyorColorMap[s] = palette[idx % palette.length]; }});
        }}

        function populateOutletDropdown(dataList) {{
            const select = document.getElementById('outletFilter');
            const currentVal = select.value;
            select.innerHTML = '<option value="ALL">-- Semua Outlet --</option>';
            const outlets = [...new Set(dataList.map(item => item.surveyee))].sort();
            outlets.forEach(o => {{
                if(o && o !== '-') {{
                    const opt = document.createElement('option');
                    opt.value = o;
                    opt.textContent = o;
                    select.appendChild(opt);
                }}
            }});
            select.value = outlets.includes(currentVal) ? currentVal : 'ALL';
        }}

        function initMap() {{
            map = L.map('map', {{ center: [-2.5489, 118.0149], zoom: 5, tap: true }});
            
            const osmStandard = L.tileLayer('https://{{s}}.tile.openstreetmap.org/{{z}}/{{x}}/{{y}}.png', {{ maxZoom: 19, attribution: '&copy; OpenStreetMap contributors' }});
            const osmDe = L.tileLayer('https://tile.openstreetmap.de/tiles/osmde/{{z}}/{{x}}/{{y}}.png', {{ maxZoom: 18, attribution: '&copy; OpenStreetMap' }});
            const esriStreet = L.tileLayer('https://server.arcgisonline.com/ArcGIS/rest/services/World_Street_Map/MapServer/tile/{{z}}/{{y}}/{{x}}', {{ maxZoom: 19, attribution: 'Tiles &copy; Esri' }});

            osmStandard.addTo(map);

            L.control.layers({{
                "OpenStreetMap (Detail Toko/Sekolah)": osmStandard,
                "OSM De": osmDe,
                "Esri Street": esriStreet
            }}).addTo(map);

            markersLayer = L.layerGroup().addTo(map);

            setTimeout(function() {{
                map.invalidateSize();
            }}, 300);

            window.addEventListener('resize', function() {{
                setTimeout(function() {{ map.invalidateSize(); }}, 200);
            }});
        }}

        function getFilteredSubset(region, asm, surveyor) {{
            let subset = surveyData;
            if (region && region !== 'ALL') {{
                subset = subset.filter(item => item.region === region);
            }}
            if (asm && asm !== 'ALL') {{
                subset = subset.filter(item => item.asm === asm);
            }}
            if (surveyor && surveyor !== 'ALL') {{
                subset = subset.filter(item => item.surveyor === surveyor);
            }}
            return subset;
        }}

        function onRegionChange() {{
            const reg = document.getElementById('regionFilter').value;
            const asm = document.getElementById('asmFilter').value;
            const srv = document.getElementById('surveyorFilter').value;

            let asmSubset = getFilteredSubset(reg, 'ALL', 'ALL');
            populateAsmDropdown(asmSubset);

            let srvSubset = getFilteredSubset(reg, asm, 'ALL');
            populateSurveyorDropdown(srvSubset);

            let outSubset = getFilteredSubset(reg, asm, srv);
            populateOutletDropdown(outSubset);

            renderDashboard(reg, document.getElementById('asmFilter').value, document.getElementById('surveyorFilter').value, 'ALL');
        }}

        function onAsmChange() {{
            const reg = document.getElementById('regionFilter').value;
            const asm = document.getElementById('asmFilter').value;
            const srv = document.getElementById('surveyorFilter').value;

            let srvSubset = getFilteredSubset(reg, asm, 'ALL');
            populateSurveyorDropdown(srvSubset);

            let outSubset = getFilteredSubset(reg, asm, srv);
            populateOutletDropdown(outSubset);

            renderDashboard(reg, asm, document.getElementById('surveyorFilter').value, 'ALL');
        }}

        function onSurveyorChange() {{
            const reg = document.getElementById('regionFilter').value;
            const asm = document.getElementById('asmFilter').value;
            const srv = document.getElementById('surveyorFilter').value;

            let outSubset = getFilteredSubset(reg, asm, srv);
            populateOutletDropdown(outSubset);

            renderDashboard(reg, asm, srv, 'ALL');
        }}

        function onOutletChange() {{
            const reg = document.getElementById('regionFilter').value;
            const asm = document.getElementById('asmFilter').value;
            const srv = document.getElementById('surveyorFilter').value;
            const out = document.getElementById('outletFilter').value;
            renderDashboard(reg, asm, srv, out);
        }}

        function renderDashboard(reg, asm, srv, out) {{
            const listContainer = document.getElementById('outletListContainer');
            listContainer.innerHTML = '';
            markersLayer.clearLayers();
            
            let filteredData = surveyData;
            if (reg !== 'ALL') filteredData = filteredData.filter(item => item.region === reg);
            if (asm !== 'ALL') filteredData = filteredData.filter(item => item.asm === asm);
            if (srv !== 'ALL') filteredData = filteredData.filter(item => item.surveyor === srv);
            if (out !== 'ALL') filteredData = filteredData.filter(item => item.surveyee === out);
            
            document.getElementById('totalOutletStat').innerText = filteredData.length;
            document.getElementById('counterBadge').innerText = filteredData.length + ' Outlet';

            let bounds = [];
            filteredData.forEach((item) => {{
                bounds.push([item.lat, item.lon]);
                const itemDiv = document.createElement('div');
                itemDiv.className = 'outlet-item';
                itemDiv.innerHTML = '<div class="outlet-name">' + item.surveyee + '</div><div class="outlet-info"><span>ID: ' + item.surveyee_id + '</span><span style="color: #2563eb; font-weight: 600;">Qty: ' + item.qty + '</span></div><div style="font-size: 10px; color: #94a3b8; margin-top: 2px;">📍 ' + item.lat + ', ' + item.lon + '</div>';
                
                let marker = createLeafletMarker(item);

                itemDiv.onclick = function() {{
                    document.querySelectorAll('.outlet-item').forEach(function(el) {{ el.classList.remove('active'); }});
                    itemDiv.classList.add('active');
                    map.setView([item.lat, item.lon], 16);
                    marker.openPopup();
                }};
                listContainer.appendChild(itemDiv);
            }});

            if (bounds.length > 0) {{
                map.fitBounds(bounds, {{ padding: [30, 30], maxZoom: 15 }});
            }}
            setTimeout(function() {{ map.invalidateSize(); }}, 100);
        }}

        function createLeafletMarker(item) {{
            const color = surveyorColorMap[item.surveyor] || '#2563eb';
            const customIcon = L.divIcon({{ className: 'custom-div-icon', html: '<div style="background-color: ' + color + '; width: 14px; height: 14px; border-radius: 50%; border: 2px solid white; box-shadow: 0 0 6px rgba(0,0,0,0.4);"></div>', iconSize: [14, 14], iconAnchor: [7, 7] }});
            
            let fotoHtml = '';
            if (item.foto_list && item.foto_list.length > 0) {{
                fotoHtml = '<div style="margin-top: 6px; text-align: center; max-height: 220px; overflow-y: auto; padding-right: 4px;">';
                item.foto_list.forEach((fPath, idx) => {{
                    fotoHtml += '<div style="margin-bottom: 6px; background: #f8fafc; padding: 4px; border-radius: 4px; border: 1px solid #e2e8f0;"><a href="' + fPath + '" target="_blank"><img src="' + fPath + '" width="220" style="border-radius: 4px; border: 1px solid #cbd5e1; display: block; margin: 0 auto;"></a><small style="color: #64748b; font-size: 10px; display: block; margin-top: 3px;">Foto ' + (idx + 1) + ' (Klik untuk memperbesar)</small></div>';
                }});
                fotoHtml += '</div>';
            }} else {{
                fotoHtml = '<div style="margin-top: 6px; color: #64748b; font-style: italic; font-size: 10px;">Tidak ada foto kunjungan.</div>';
            }}
            
            const popupContent = '<b style="color: #0f172a; font-size: 13px;">' + item.surveyee + '</b><br><span style="color: #64748b;">ID:</span> <b>' + item.surveyee_id + '</b><br><hr style="margin: 4px 0; border: 0; border-top: 1px solid #cbd5e1;"><span style="color: #64748b;">Waktu:</span> ' + item.timestamp + '<br><span style="color: #64748b;">Surveyor:</span> ' + item.surveyor + '<br><span style="color: #64748b;">Region/ASM:</span> ' + item.region + ' / ' + item.asm + '<br><span style="color: #64748b;">Qty Total:</span> <span style="color: #2563eb; font-weight: bold;">' + item.qty + '</span><br><span style="color: #64748b;">GPS:</span> <code style="background: #f1f5f9; padding: 1px 3px; border-radius: 3px; font-size: 10px;">' + item.lat + ', ' + item.lon + '</code>' + fotoHtml;
            
            const marker = L.marker([item.lat, item.lon], {{ icon: customIcon }}).bindPopup(popupContent);
            markersLayer.addLayer(marker);
            return marker;
        }}

        window.onload = initDashboard;
    </script>
</body>
</html>
"""

        with open(FILE_HTML_PETA, 'w', encoding='utf-8') as f_html:
            f_html.write(multilayer_dashboard_html)

        print(f"\n[RINGKASAN PROSES FINAL]")
        print(f"1. Total baris data diproses                : {len(list_log_detail)} baris")
        print(f"2. Total baris digabung/akumulasi Qty         : {total_baris_digabung} baris")
        print(f"3. Total outlet unik tersimpan di Excel Utama: {len(df_clean)} baris")
        print(f"4. Total foto unik yang berhasil didownload    : {total_foto_diekstrak} foto")
        print(f"5. Total titik GPS valid di Peta GIS          : {len(list_data_peta)} titik")
        print(f"6. File Excel Utama disimpan di               : '{path_excel_output}'")
        print(f"7. File Log Detail disimpan di                : '{FILE_LOG}'")
        print(f"8. File Peta HTML (Mobile-Responsive GIS)      : '{FILE_HTML_PETA}'")
        
        # --- FITUR OTOMATIS UPLOAD KE GITHUB (DENGAN PATH ABSOLUT GIT & AUTO .GITIGNORE) ---
        print("\n-> Sedang melakukan Auto-Sync & Push ke GitHub...")
        try:
            waktu_sekarang = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            
            # Otomatis pastikan folder source/ diabaikan oleh git agar tidak terkena limit ukuran file GitHub
            gitignore_path = os.path.join(BASE_DIR, ".gitignore")
            ignore_content = ""
            if os.path.exists(gitignore_path):
                with open(gitignore_path, "r", encoding="utf-8") as gf:
                    ignore_content = gf.read()
            if "source/" not in ignore_content:
                with open(gitignore_path, "a", encoding="utf-8") as gf:
                    gf.write("\nsource/\n")

            # Path absolut Git di Windows (fallback ke 'git' jika tidak ditemukan di path default)
            git_path = r"C:\Program Files\Git\cmd\git.exe"
            if not os.path.exists(git_path):
                git_path = "git"

            subprocess.run(f'"{git_path}" add .', cwd=BASE_DIR, shell=True, check=True)
            subprocess.run(f'"{git_path}" commit -m "Auto-update data surveyor & peta pada {waktu_sekarang}"', cwd=BASE_DIR, shell=True, check=True)
            subprocess.run(f'"{git_path}" push', cwd=BASE_DIR, shell=True, check=True)
            print("[SUKSES] Perubahan berhasil di-push otomatis ke GitHub! 🎉")
        except subprocess.CalledProcessError as git_err:
            print(f"[WARNING] Gagal auto-push ke GitHub: {git_err}")
            print("Tips: Pastikan koneksi internet stabil atau lakukan 'git push' manual jika diperlukan.")
        # -----------------------------------------------------------------
        
    except PermissionError as e:
        print(f"\n[ERROR] Gagal menyimpan file karena sedang dibuka di Microsoft Excel! ({e})")
else:
    print("[WARNING] Total baris gabungan kosong.")

print("\n=== [SELESAI ALL PROCESS] == Mobile-Ready Dashboard Generated ==")